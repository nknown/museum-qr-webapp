import os
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from pathlib import Path

from django.contrib.auth.models import User
from django.core.files.base import ContentFile
from django.db import transaction
from django.utils.text import slugify
from wagtail.images import get_image_model
from wagtail.models import Page

from main_app.models import (
    Exhibition,
    ExhibitionOnSitePage,
    ExhibitionPage,
    ExhibitionPiecesListPage,
    ExhibitionUserDevicePage,
    Piece,
    PieceOnSitePage,
    PiecePage,
    PieceUserDevicePage,
    SlideshowSlidePage,
)

try:
    import ezodf
except ImportError:
    ezodf = None

try:
    import openpyxl
except ImportError:
    openpyxl = None


class FeederImportError(Exception):
    pass


@dataclass
class FeederImportResult:
    exhibition: Exhibition
    created: int = 0
    updated: int = 0
    ignored: int = 0
    attributes: int = 0


class FeederImporter:
    supported_extensions = ('.ods', '.xlsx', '.xlsm')

    def __init__(self, writer=None):
        self.writer = writer or (lambda message: None)
        self.file_path = None

    def import_path(self, file_path, exhibition=None, admin_user=None, assign_superusers=False):
        self.file_path = str(file_path)
        if not os.path.exists(self.file_path):
            raise FeederImportError(f'File not found: {self.file_path}')

        extension = Path(self.file_path).suffix.lower()
        if extension not in self.supported_extensions:
            raise FeederImportError('Upload an ODS or Excel .xlsx file.')

        images, attributes, pieces_data = self._parse_file(extension)
        return self.import_data(
            Path(self.file_path).stem.strip(),
            images,
            attributes,
            pieces_data,
            exhibition=exhibition,
            admin_user=admin_user,
            assign_superusers=assign_superusers,
        )

    def import_uploaded_file(self, uploaded_file, exhibition, admin_user=None, assign_superusers=False):
        suffix = Path(uploaded_file.name).suffix.lower()
        if suffix not in self.supported_extensions:
            raise FeederImportError('Upload an ODS or Excel .xlsx file.')

        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            for chunk in uploaded_file.chunks():
                tmp.write(chunk)
            tmp_path = tmp.name

        try:
            self.file_path = tmp_path
            images, attributes, pieces_data = self._parse_file(suffix)
            return self.import_data(
                Path(uploaded_file.name).stem.strip(),
                images,
                attributes,
                pieces_data,
                exhibition=exhibition,
                admin_user=admin_user,
                assign_superusers=assign_superusers,
            )
        finally:
            try:
                os.unlink(tmp_path)
            except OSError:
                pass

    def import_data(
        self,
        exhibition_name,
        images,
        attributes,
        pieces_data,
        exhibition=None,
        admin_user=None,
        assign_superusers=False,
    ):
        if not attributes:
            raise FeederImportError('No attributes found in spreadsheet.')
        if not pieces_data:
            raise FeederImportError('No piece data found in spreadsheet.')

        from main_app.access import can_curate, is_curator, curator_museum, managed_exhibitions
        museum = None
        candidates = Exhibition.objects.all()
        if admin_user is not None:
            if not can_curate(admin_user):
                raise FeederImportError('Only administrators and curators can import content.')
            candidates = managed_exhibitions(admin_user)
            if is_curator(admin_user):
                try:
                    museum = curator_museum(admin_user)
                except Exception as exc:
                    raise FeederImportError(str(exc)) from exc
            if exhibition is not None and not candidates.filter(pk=exhibition.pk).exists():
                raise FeederImportError('You cannot import into this exhibition.')
        with transaction.atomic():
            if exhibition is None:
                exhibition = candidates.filter(title=exhibition_name).order_by('pk').first()
                created_exhibition = exhibition is None
                if exhibition is None:
                    exhibition = Exhibition.objects.create(
                        title=exhibition_name,
                        description='Auto-imported from feeder',
                        museum=museum,
                    )
                self._assign_exhibition_admins(
                    exhibition,
                    admin_user=admin_user,
                    assign_superusers=assign_superusers,
                    only_if_empty=False,
                )
            elif slugify(exhibition.title) != slugify(exhibition_name):
                raise FeederImportError(
                    f'Spreadsheet name "{exhibition_name}" does not match exhibition "{exhibition.title}".'
                )
            else:
                self._assign_exhibition_admins(
                    exhibition,
                    admin_user=admin_user,
                    assign_superusers=assign_superusers,
                    only_if_empty=False,
                )

            self._ensure_exhibition_slideshow_tree(exhibition)
            self._update_exhibition_attributes(exhibition, attributes)
            result = self._upsert_pieces(exhibition, exhibition.pieces_attributes, pieces_data, images)
            result.attributes = len(exhibition.pieces_attributes)
            return result

    def _assign_exhibition_admins(self, exhibition, admin_user=None, assign_superusers=False, only_if_empty=False):
        if only_if_empty and exhibition.admins.exists():
            return

        if admin_user:
            exhibition.admins.add(admin_user)
            return

        if assign_superusers:
            superusers = User.objects.filter(is_superuser=True)
            if superusers.exists():
                exhibition.admins.add(*superusers)

    def _parse_file(self, extension):
        if extension == '.ods':
            return self._parse_ods()
        return self._parse_xlsx()

    def _parse_ods(self):
        if ezodf is None:
            raise FeederImportError('ezodf is not installed. Run: pip install ezodf lxml')

        try:
            doc = ezodf.opendoc(self.file_path)
            sheet = doc.sheets[0]
        except Exception as exc:
            raise FeederImportError(f'Error reading spreadsheet: {exc}')

        image_refs = self._get_image_refs_from_ods_xml(max_col=None)
        attributes = []
        pieces_data = {}
        max_row = sheet.nrows()
        max_col = sheet.ncols()

        if not image_refs:
            image_refs = self._get_image_refs_from_ods_xml(max_col=max_col)

        if max_row < 2 or max_col < 2:
            raise FeederImportError('Spreadsheet must have at least 2 rows and 2 columns.')

        if not image_refs:
            for col_idx in range(1, max_col):
                image_filename = self._get_image_filename_from_cell(sheet[(0, col_idx)])
                if image_filename:
                    image_refs[col_idx] = image_filename

        images = self._extract_images_from_ods(image_refs)

        for row_idx in range(1, max_row):
            value = self._get_ods_cell_value(sheet[(row_idx, 0)])
            if value:
                attributes.append(value)

        for col_idx in range(1, max_col):
            piece_values = {}
            for row_idx, attr_name in enumerate(attributes, start=1):
                value = self._get_ods_cell_value(sheet[(row_idx, col_idx)])
                piece_values[attr_name.lower()] = value or ''
            pieces_data[col_idx] = {'attributes': piece_values, 'image_col': col_idx}

        return images, attributes, pieces_data

    def _parse_xlsx(self):
        if openpyxl is None:
            raise FeederImportError('openpyxl is not installed. Run: pip install openpyxl')

        try:
            workbook = openpyxl.load_workbook(self.file_path)
            sheet = workbook.worksheets[0]
        except Exception as exc:
            raise FeederImportError(f'Error reading spreadsheet: {exc}')

        max_row = sheet.max_row or 0
        max_col = sheet.max_column or 0
        if max_row < 2 or max_col < 2:
            raise FeederImportError('Spreadsheet must have at least 2 rows and 2 columns.')

        attributes = []
        for row_idx in range(2, max_row + 1):
            value = self._clean_cell_value(sheet.cell(row=row_idx, column=1).value)
            if value:
                attributes.append(value)

        pieces_data = {}
        for col_idx in range(2, max_col + 1):
            piece_values = {}
            for row_offset, attr_name in enumerate(attributes, start=2):
                value = self._clean_cell_value(sheet.cell(row=row_offset, column=col_idx).value)
                piece_values[attr_name.lower()] = value or ''
            pieces_data[col_idx - 1] = {'attributes': piece_values, 'image_col': col_idx - 1}

        return self._extract_images_from_xlsx(sheet), attributes, pieces_data

    def _get_image_refs_from_ods_xml(self, max_col=None):
        image_refs = {}
        if not self.file_path:
            return image_refs

        try:
            with zipfile.ZipFile(self.file_path, 'r') as ods_zip:
                content_xml = ods_zip.read('content.xml')
        except Exception:
            return image_refs

        namespaces = {
            'table': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0',
            'draw': 'urn:oasis:names:tc:opendocument:xmlns:drawing:1.0',
            'xlink': 'http://www.w3.org/1999/xlink',
        }

        try:
            root = ET.fromstring(content_xml)
            table = root.find('.//table:table', namespaces)
            first_row = table.find('table:table-row', namespaces) if table is not None else None
            if first_row is None:
                return image_refs

            col_idx = 0
            repeat_attr = f"{{{namespaces['table']}}}number-columns-repeated"
            href_attr = f"{{{namespaces['xlink']}}}href"
            for cell in first_row.findall('table:table-cell', namespaces):
                repeat = int(cell.get(repeat_attr, '1'))
                image_filenames = [
                    image.get(href_attr)
                    for image in cell.findall('.//draw:image', namespaces)
                    if image.get(href_attr)
                ]
                if image_filenames and col_idx > 0 and (max_col is None or col_idx < max_col):
                    image_refs[col_idx] = image_filenames[-1]
                col_idx += repeat
                if max_col is not None and col_idx >= max_col:
                    break
        except Exception:
            return image_refs

        return image_refs

    def _extract_images_from_ods(self, image_refs):
        images = {}
        try:
            with zipfile.ZipFile(self.file_path, 'r') as ods_zip:
                for col_idx, image_ref in image_refs.items():
                    try:
                        images[col_idx] = ods_zip.read(image_ref)
                    except KeyError:
                        self.writer(f'Warning: Could not find image {image_ref}')
        except Exception as exc:
            self.writer(f'Warning: Could not extract images: {exc}')
        return images

    def _extract_images_from_xlsx(self, sheet):
        images = {}
        for image in getattr(sheet, '_images', []):
            anchor = getattr(image, 'anchor', None)
            marker = getattr(anchor, '_from', None)
            if marker is None or marker.row != 0 or marker.col < 1:
                continue
            try:
                images[marker.col] = image._data()
            except Exception as exc:
                self.writer(f'Warning: Could not extract an Excel image: {exc}')
        return images

    def _get_ods_cell_value(self, cell):
        if cell is None:
            return ''
        try:
            return self._clean_cell_value(cell.plaintext() if hasattr(cell, 'plaintext') else cell)
        except Exception:
            return ''

    def _get_image_filename_from_cell(self, cell):
        try:
            if hasattr(cell, 'node') and cell.node is not None:
                for child in cell.node:
                    if 'image' in child.tag.lower():
                        href = child.get('{http://www.w3.org/1999/xlink}href')
                        if href:
                            return href
        except Exception:
            return None
        return None

    def _clean_cell_value(self, value):
        if value is None:
            return ''
        return str(value).strip()

    def _update_exhibition_attributes(self, exhibition, attributes):
        normalized = []
        for attr in attributes:
            attr = str(attr).strip()
            if attr and attr.lower() not in [existing.lower() for existing in normalized]:
                normalized.append(attr)

        existing = exhibition.pieces_attributes if isinstance(exhibition.pieces_attributes, list) else []
        merged = list(existing)
        for attr in normalized:
            if attr.lower() not in [existing_attr.lower() for existing_attr in merged]:
                merged.append(attr)

        if not any(attr.lower() == 'name' for attr in merged):
            merged.insert(0, 'Name')

        exhibition.pieces_attributes = merged
        exhibition.save(update_fields=['pieces_attributes'])

    def _ensure_exhibition_slideshow_tree(self, exhibition):
        from main_app.views.models_views import _get_or_create_child_page, _save_as_draft, _unique_child_slug

        if exhibition.slideshow_page_id:
            exhibition_page = exhibition.slideshow_page.specific
        else:
            root = Page.get_first_root_node()
            exhibition_page = ExhibitionPage(
                title=exhibition.title,
                slug=_unique_child_slug(root, exhibition.title),
                description=exhibition.description,
            )
            root.add_child(instance=exhibition_page)
            _save_as_draft(exhibition_page)
            exhibition.slideshow_page = exhibition_page
            exhibition.save(update_fields=['slideshow_page'])

        _get_or_create_child_page(exhibition_page, ExhibitionOnSitePage, title='On Site', slug='on_site')
        user_device_page = _get_or_create_child_page(
            exhibition_page,
            ExhibitionUserDevicePage,
            title='User Device',
            slug='user_device',
        )
        _get_or_create_child_page(user_device_page, SlideshowSlidePage, title='Slide 1', slug='slide_1')
        _get_or_create_child_page(exhibition_page, ExhibitionPiecesListPage, title='Pieces', slug='pieces')
        return exhibition_page

    def _upsert_pieces(self, exhibition, attributes, pieces_data, images):
        result = FeederImportResult(exhibition=exhibition)
        pieces_list_page = exhibition.slideshow_page.get_children().filter(slug='pieces').first()
        if pieces_list_page is None:
            raise FeederImportError('Could not find the exhibition Pieces page.')
        pieces_list_page = pieces_list_page.specific

        for col_idx, data in pieces_data.items():
            piece_attrs = data['attributes']
            piece_name = piece_attrs.get('name', '').strip()
            if not piece_name:
                result.ignored += 1
                continue

            piece = self._find_piece_by_name(exhibition, piece_name)
            existing_values = {}
            if piece and isinstance(piece.attributes, list):
                existing_values = {
                    str(item.get('name', '')).strip().lower(): str(item.get('value', '')).strip()
                    for item in piece.attributes
                    if isinstance(item, dict)
                }

            attributes_payload = []
            for attr in attributes:
                key = attr.lower()
                attributes_payload.append({
                    'name': attr,
                    'value': piece_attrs[key] if key in piece_attrs else existing_values.get(key, ''),
                })

            if piece:
                changed = False
                if piece.attributes != attributes_payload:
                    piece.attributes = attributes_payload
                    changed = True
                if not piece.slideshow_page_id:
                    piece.slideshow_page = self._create_piece_page(pieces_list_page, piece, piece_name)
                    changed = True
                if changed:
                    piece.save()
                    result.updated += 1
                else:
                    result.ignored += 1
            else:
                piece = Piece.objects.create(exhibition=exhibition, attributes=attributes_payload)
                piece.slideshow_page = self._create_piece_page(pieces_list_page, piece, piece_name)
                piece.save(update_fields=['slideshow_page'])
                result.created += 1

            if col_idx in images and images[col_idx]:
                self._assign_piece_image(piece, piece_name, images[col_idx])

        return result

    def _find_piece_by_name(self, exhibition, piece_name):
        wanted = piece_name.strip().lower()
        for piece in exhibition.pieces.all():
            if str(piece).strip().lower() == wanted:
                return piece
        return None

    def _create_piece_page(self, pieces_list_page, piece, piece_name):
        from main_app.views.models_views import _get_or_create_child_page, _save_as_draft, _unique_child_slug

        piece_page = PiecePage(
            title=piece_name,
            slug=_unique_child_slug(pieces_list_page, f'piece-{piece.pk}-{piece_name}'),
        )
        pieces_list_page.add_child(instance=piece_page)
        _save_as_draft(piece_page)
        _get_or_create_child_page(piece_page, PieceOnSitePage, title='On Site', slug='on_site')
        user_device_page = _get_or_create_child_page(piece_page, PieceUserDevicePage, title='User Device', slug='user_device')
        _get_or_create_child_page(user_device_page, SlideshowSlidePage, title='Slide 1', slug='slide_1')
        return piece_page

    def _assign_piece_image(self, piece, piece_name, image_data):
        from main_app.views.models_views import _ensure_presentation_folders

        wagtail_image = self._save_image_to_wagtail(piece_name, image_data)
        if not wagtail_image or not piece.slideshow_page_id:
            return

        folders = _ensure_presentation_folders(piece.slideshow_page)
        for page in folders.values():
            if page and getattr(page, 'hero_image_id', None) != wagtail_image.pk:
                page.hero_image = wagtail_image
                page.save()

    def _save_image_to_wagtail(self, name, image_data):
        extension = self._get_image_extension(image_data) or 'jpg'
        Image = get_image_model()
        return Image.objects.create(
            title=name,
            file=ContentFile(image_data, name=f'{slugify(name) or "piece"}.{extension}'),
        )

    def _get_image_extension(self, image_data):
        if image_data.startswith(b'\xff\xd8\xff'):
            return 'jpg'
        if image_data.startswith(b'\x89PNG'):
            return 'png'
        if image_data.startswith(b'GIF8'):
            return 'gif'
        if image_data.startswith(b'RIFF') and b'WEBP' in image_data[:20]:
            return 'webp'
        return None
